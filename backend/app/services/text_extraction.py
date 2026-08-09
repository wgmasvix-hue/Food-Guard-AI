"""Best-effort text extraction for uploaded documents, so files a user
drops into Documents (PDFs, Word docs, plain text) become searchable by
the RAG pipeline (see app.services.search) instead of only documents
created with an explicit content_text field.

Extraction is deliberately forgiving: any failure (corrupt file,
unsupported/encrypted PDF, missing dependency) just yields no text
rather than blocking the upload — the file is still stored and
downloadable, it simply won't surface in search/AI context.
"""
import io
import logging

logger = logging.getLogger(__name__)

# Cap extracted text so a huge upload can't blow up prompt sizes or the
# tsvector index; RAG only ever needs the gist of a document, not every
# page verbatim.
MAX_EXTRACTED_CHARS = 50_000

_PDF_CONTENT_TYPES = {"application/pdf"}
_DOCX_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
_TEXT_CONTENT_TYPES = {"text/plain", "text/markdown", "text/csv"}


def _by_extension(filename: str | None) -> str | None:
    if not filename or "." not in filename:
        return None
    return filename.rsplit(".", 1)[-1].lower()


def _extract_pdf(data: bytes) -> str | None:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency always installed in prod
        logger.warning("pypdf not installed; skipping PDF text extraction")
        return None
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages).strip() or None
    except Exception:
        logger.warning("Failed to extract text from PDF upload", exc_info=True)
        return None


def _extract_docx(data: bytes) -> str | None:
    try:
        import docx
    except ImportError:  # pragma: no cover - dependency always installed in prod
        logger.warning("python-docx not installed; skipping DOCX text extraction")
        return None
    try:
        document = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(p for p in paragraphs if p.strip()).strip() or None
    except Exception:
        logger.warning("Failed to extract text from DOCX upload", exc_info=True)
        return None


def _extract_plain_text(data: bytes) -> str | None:
    try:
        return data.decode("utf-8").strip() or None
    except UnicodeDecodeError:
        try:
            return data.decode("latin-1").strip() or None
        except Exception:
            return None


def extract_text(data: bytes, content_type: str | None, filename: str | None) -> str | None:
    """Return best-effort plain text extracted from an uploaded file, or
    None if the format isn't supported / extraction failed."""
    ext = _by_extension(filename)
    text: str | None = None

    if content_type in _PDF_CONTENT_TYPES or ext == "pdf":
        text = _extract_pdf(data)
    elif content_type in _DOCX_CONTENT_TYPES or ext == "docx":
        text = _extract_docx(data)
    elif content_type in _TEXT_CONTENT_TYPES or ext in {"txt", "md", "csv"}:
        text = _extract_plain_text(data)

    if text and len(text) > MAX_EXTRACTED_CHARS:
        text = text[:MAX_EXTRACTED_CHARS]
    return text
